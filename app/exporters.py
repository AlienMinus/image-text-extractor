import os
import uuid
import io
from docx import Document
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from gtts import gTTS
from xml.sax.saxutils import escape


def export_docx(text):
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_xlsx(text):
    buffer = io.BytesIO()
    wb = Workbook()
    ws = wb.active

    for i, line in enumerate(text.split("\n"), 1):
        ws[f"A{i}"] = line

    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_pdf(text):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)
    styles = getSampleStyleSheet()
    elements = []

    for line in text.split("\n"):
        # Escape text to prevent XML parsing errors in ReportLab
        elements.append(Paragraph(escape(line), styles["Normal"]))
        elements.append(Spacer(1, 0.2 * inch))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def export_md(text):
    buffer = io.BytesIO()
    buffer.write(text.encode("utf-8"))
    buffer.seek(0)
    return buffer


def export_mp3(text, language='en'):
    buffer = io.BytesIO()
    
    # Map Tesseract 3-letter codes to 2-letter ISO codes
    lang_map = {
        'eng': 'en', 'spa': 'es', 'fra': 'fr', 'deu': 'de', 'ita': 'it',
        'auto': 'en'
    }
    
    lang = lang_map.get(language, language)
    
    try:
        tts = gTTS(text=text, lang=lang)
        tts.write_to_fp(buffer)
    except ValueError:
        # Fallback to English if language is not supported
        tts = gTTS(text=text, lang='en')
        tts.write_to_fp(buffer)
        
    buffer.seek(0)
    return buffer